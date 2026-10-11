#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gera_river_docx.py — atril de River Flows in You (PoC 5), versão papel.

Importa gera_river.py (que parseia o .mscz na hora): FONTE ÚNICA de verdade —
nenhuma nota é retranscrita aqui. O que este arquivo acrescenta é apenas a
roupagem de papel: um DOCX A4 pensado pro ATRIL (imprimir e pôr na estante)
e pra ANOTAÇÃO À MÃO (o dedilhado que o aluno validar na prática).

Saída: ../saidas/atril-river-flows-in-you.docx
Roda:  python3 05-river-flows-in-you/trabalho/gera_river_docx.py
Requer: pip install python-docx pillow
"""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import gera_river as gr
from gera_river import (S1, S2, RIO_MAPA, SECOES, RIO_CELULAS, RIO_GRAVES,
                        PADROES_MD, FASES, COSTURAS, dur_letra, nome)

# ───────────────────────────────── cores (o rio em papel: fundo branco, imprime barato)
INK   = RGBColor(0x1A, 0x32, 0x3C)   # texto principal (azul-noite)
DARK  = RGBColor(0x0E, 0x50, 0x63)   # títulos (petróleo)
TEAL  = RGBColor(0x1D, 0x6F, 0x8C)   # mão direita — melodia (ciano em papel)
PEDAL = RGBColor(0x6B, 0x8F, 0x99)   # o pedal A4 e o grave das díades
GOLD  = RGBColor(0x8A, 0x6D, 0x2F)   # rótulos / células do rio
GOLDT = RGBColor(0x9A, 0x64, 0x0C)   # notas da mão esquerda
PEAK  = RGBColor(0xC0, 0x39, 0x2B)   # o teto E6 (a cascata)
MUTED = RGBColor(0x64, 0x7A, 0x84)   # rótulos pequenos / durações
NOTE  = RGBColor(0x4A, 0x5A, 0x6B)   # observações

CREAM = "EAF6FA"   # caixa-dica (água rasa em papel)
GOLDB = "FFF3D6"   # caixa dourada (o rio)

# ───────────────────────────────── helpers docx
def shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def box(doc, title, text, fill=CREAM, icon="💡", tcol=DARK):
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_before = Pt(8)
    p1.paragraph_format.space_after = Pt(0)
    r = p1.add_run(f"{icon} {title}")
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = tcol
    shade(p1, fill)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)
    r2.font.color.rgb = NOTE
    shade(p2, fill)
    return p2


def sub(doc, text, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(size)
    r.font.color.rgb = NOTE
    p.paragraph_format.space_after = Pt(4)
    return p


def bullets(doc, items, size=10):
    for b in items:
        p = doc.add_paragraph(b, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(size)


def numlist(doc, items):
    for lead, rest in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(lead)
        r1.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = INK
        r2 = p.add_run(rest)
        r2.font.size = Pt(10)
        r2.font.color.rgb = NOTE


def classe_nota(pitch, mao, baixo_da_díade=False):
    if mao == 'E':
        return GOLDT
    if pitch == 88:                    # E6 — o teto
        return PEAK
    if pitch == 69 or baixo_da_díade:  # A4 — o polegar mora aqui
        return PEDAL
    return TEAL


# ───────────────────────────────── compasso → tabela
def add_secao(doc, nom, i0, i1, titulo, desc):
    doc.add_heading(f"{nom} · m{i0}–{i1} · {titulo}", level=2)
    sub(doc, desc, 9)
    n = i1 - i0 + 1
    ncols = max(len(m['voices'][0]) for m in S1[i0-1:i1] + S2[i0-1:i1])
    tbl = doc.add_table(rows=2 * n, cols=ncols + 1)
    tbl.style = "Table Grid"

    LARGO, USO = Inches(0.95), 6.05
    wcol = Inches(USO / ncols)
    for k in range(n):
        mi = i0 + k
        evs_d = S1[mi-1]['voices'][0]
        evs_e = S2[mi-1]['voices'][0]
        # célula-rótulo (2 linhas unidas): nº + rio
        lab = tbl.cell(2*k, 0).merge(tbl.cell(2*k+1, 0))
        lp = lab.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        r = lp.add_run(f"m{mi}")
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = INK
        rio = ' · '.join(RIO_MAPA.get(mi, []))
        if rio:
            lp2 = lab.add_paragraph()
            lp2.paragraph_format.space_after = Pt(0)
            rr = lp2.add_run(rio)
            rr.font.size = Pt(7)
            rr.font.color.rgb = GOLD
        for lin, (evs, mao) in enumerate(((evs_d, 'D'), (evs_e, 'E'))):
            for ci in range(ncols + 1):
                tbl.cell(2*k+lin, ci).width = LARGO if ci == 0 else wcol
            for ci, ev in enumerate(evs, start=1):
                cell = tbl.cell(2*k+lin, ci)
                cell.width = wcol
                p1 = cell.paragraphs[0]
                p1.paragraph_format.space_after = Pt(0)
                p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p2 = cell.add_paragraph()
                p2.paragraph_format.space_after = Pt(0)
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if ev['t'] == 'gap':
                    r1 = p1.add_run("—"); r1.font.size = Pt(8); r1.font.color.rgb = MUTED
                    r2 = p2.add_run(" "); r2.font.size = Pt(6)
                    continue
                if ev['t'] == 'rest':
                    r1 = p1.add_run("·"); r1.font.size = Pt(8); r1.font.color.rgb = MUTED
                    r2 = p2.add_run(dur_letra(ev) + "R"); r2.font.size = Pt(6.5); r2.font.color.rgb = MUTED
                    continue
                notas = sorted(ev['notas'], key=lambda nn: nn['pitch'])
                for j, nn in enumerate(notas):
                    if j:
                        sep = p1.add_run("+")
                        sep.font.size = Pt(7)
                        sep.font.color.rgb = MUTED
                    r1 = p1.add_run(("~" if ev.get('grace') else "") + nome(nn['pitch'], nn['tpc']))
                    r1.bold = True
                    r1.font.size = Pt(8)
                    r1.font.name = "Consolas"
                    r1.font.color.rgb = classe_nota(nn['pitch'], mao,
                                                     baixo_da_díade=(len(notas) > 1 and j == 0))
                r2 = p2.add_run("~" if ev.get('grace') else dur_letra(ev))
                r2.font.size = Pt(6.5)
                r2.font.color.rgb = MUTED


# ───────────────────────────────── teclado-guia (PIL)
def draw_keyboard(path):
    from PIL import Image, ImageDraw, ImageFont
    import os
    LH = {n['pitch'] for n in gr.notas_mao(S2)}
    RH = {n['pitch'] for n in gr.notas_mao(S1)}

    W, H, BW, BH, M = 44, 140, 26, 88, 18
    lo, hi = 40, 88                                    # E2..E6
    whites = [p for p in range(lo, hi + 1) if p % 12 in (0, 2, 4, 5, 7, 9, 11)]
    blacks = [p for p in range(lo, hi + 1) if p % 12 in (1, 3, 6, 8, 10)]
    y0 = 92
    IMGW = 2 * M + len(whites) * W
    img = Image.new("RGB", (IMGW, y0 + H + 66), "#ffffff")
    d = ImageDraw.Draw(img)

    def font(sz, bold=True):
        p = ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold
             else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        return ImageFont.truetype(p, sz) if os.path.exists(p) else ImageFont.load_default()

    f_key, f_leg = font(15), font(15)

    def center(cx, cy, txt, f, col):
        bb = d.textbbox((0, 0), txt, font=f)
        d.text((cx - (bb[2] - bb[0]) / 2, cy - (bb[3] - bb[1]) / 2), txt, font=f, fill=col)

    x_pos, x = {}, M
    for p in whites:
        if p in LH:      fill, tcol = "#F5C64B", "#4a3a10"
        elif p in RH:    fill, tcol = "#9FD8E8", "#0e3d4d"
        else:            fill, tcol = "#F7F3E8", None
        d.rectangle([x, y0, x + W - 2, y0 + H - 1], fill=fill, outline="#8d9aa8", width=2)
        if tcol:
            center(x + (W - 2) / 2, y0 + H - 24, nome(p), f_key, tcol)
        if p == 69:                                      # A4 — o ponto de encontro
            cx = x + (W - 2) / 2
            d.ellipse([cx - 9, y0 + 10, cx + 9, y0 + 28], outline="#e67e22", width=3)
        x_pos[p] = x
        x += W
    for p in blacks:
        bx = x_pos[p - 1] + W - BW // 2 - 1
        if p in LH:      fill, tcol = "#D9A935", "#241505"
        elif p in RH:    fill, tcol = "#4FA8C4", "#06222c"
        else:            fill, tcol = "#23262E", None
        d.rectangle([bx, y0, bx + BW - 1, y0 + BH - 1], fill=fill, outline="#14171b", width=2)
        if tcol:
            center(bx + BW / 2, y0 + BH - 14, nome(p), font(11), tcol)
        x_pos[p] = bx

    def cx(p):
        return x_pos[p] + ((W - 2) / 2 if p in whites else BW / 2)

    def callout(p, txt, ty, col, align="mid"):
        f = font(15)
        bb = d.textbbox((0, 0), txt, font=f)
        tw = bb[2] - bb[0]
        c = cx(p)
        tx = c - tw / 2 if align == "mid" else IMGW - M - tw
        d.line([c, ty + 14, c, y0 - 3], fill=col, width=1)
        d.text((tx, ty), txt, font=f, fill=col)

    callout(81, "▲ A5 — a melodia mora aqui", 16, "#1D6F8C")
    callout(88, "▲ o teto — E6 (a cascata)", 48, "#C0392B", align="right")
    callout(69, "A4: o polegar mora aqui", 16, "#e67e22")
    # E2 abaixo (a tecla está na borda esquerda)
    yb = y0 + H + 14
    d.line([cx(40), yb, cx(40) + 16, yb + 10], fill="#8A6D2F", width=1)
    d.text((cx(40) + 22, yb + 4), "▼ E2 — o fundo do rio (as seções B descem aqui)",
           font=f_leg, fill="#8A6D2F")

    ly = y0 + H + 40
    x = M
    for fill, lab in (("#F5C64B", "mão esquerda — o rio"), ("#9FD8E8", "mão direita")):
        d.rectangle([x, ly, x + 26, ly + 16], fill=fill, outline="#8d9aa8")
        d.text((x + 32, ly), lab, font=f_leg, fill="#33465a")
        x += 32 + d.textlength(lab, font=f_leg) + 24
    d.ellipse([x + 4, ly + 1, x + 20, ly + 17], outline="#e67e22", width=3)
    d.text((x + 26, ly), "A4 = o ponto de encontro", font=f_leg, fill="#33465a")
    img.save(path)


# ───────────────────────────────── montagem
def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.6)

    st = doc.styles["Normal"]
    st.font.name = "Segoe UI"
    st.font.size = Pt(11)
    st.font.color.rgb = INK
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Segoe UI"; h1.font.size = Pt(14); h1.font.bold = True; h1.font.color.rgb = DARK
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Segoe UI"; h2.font.size = Pt(12); h2.font.bold = True; h2.font.color.rgb = GOLD

    # ---------- cabeçalho
    doc.add_heading("🌊 River Flows in You — o atril", level=0)
    sub(doc, "Yiruma, 2001 · o mapa do aluno: do primeiro contato à peça inteira, as duas mãos", 11)
    badges = doc.add_paragraph()
    rb = badges.add_run("Tom: LÁ MAIOR — o original (tocar junto com o Yiruma funciona)   ·   4/4 · "
                        "escada de estudo 60→65→70→78   ·   47 compassos · 6 seções · 854 notas   ·   "
                        "sem oitavas simultâneas (mão pequena bem-vinda)   ·   9 MIDIs de estudo "
                        "acompanham este papel")
    rb.font.size = Pt(9)
    rb.font.color.rgb = MUTED

    # ---------- a história
    doc.add_heading("💡 A peça é um rio", level=1)
    for txt in (
        "Duas mãos, dois papéis que trocam de lugar: a esquerda é o RIO — quatro células que "
        "giram num loop único (F#m → D → A → E, a harmonia da peça INTEIRA, sem uma modulação); "
        "a direita é primeiro o tema, depois vira a correnteza (a onda e a cascata), e no fim "
        "devolve o convite do começo.",
        "O final é a lição mais bonita: a peça TERMINA SUSPENSA — B4 sobre E4, o acorde de "
        "dominante, sem resolver. Ela não fecha: flutua. Deixe o acorde soar e resista à "
        "tentação de tocar um Lá final “pra fechar”.",
    ):
        p = doc.add_paragraph(txt)
        p.paragraph_format.space_after = Pt(6)

    # ---------- como ler
    doc.add_heading("📖 Como ler este atril (sem ler partitura)", level=1)
    bullets(doc, [
        "Cada compasso tem 2 linhas: MD (mão direita, azul) e ME (mão esquerda, DOURADA — o rio).",
        "Cada quadradinho = uma coisa que toca, na ordem: nota em cima, duração embaixo "
        "(e=colcheia · s=semicolcheia · t=fusa · q=semínima · h=mínima · w=semibreve · ·=pausa).",
        "Notas com ~ = de graça: rapidinha, esmagada antes da seguinte (nos MIDIs elas viram um "
        "“crush” de estudo).",
        "Dourado no rótulo do compasso = as células do rio (F#m · D · A · E; ⁻ = versão funda, "
        "com o baixo uma oitava abaixo).",
        "VERMELHO = o teto E6 (o fim da cascata) · azul-petróleo = o pedal A4, onde o polegar "
        "da direita mora quase a peça inteira.",
    ])

    # ---------- teclado
    doc.add_heading("⌨️ O teclado da peça", level=1)
    try:
        draw_keyboard(SAIDAS / "teclado-river.png")
        doc.add_picture(str(SAIDAS / "teclado-river.png"), width=Inches(6.9))
        doc.paragraphs[-1].alignment = 1
    except Exception as e:
        p = doc.add_paragraph(f"(Teclado-guia indisponível nesta geração: {e})")
        p.runs[0].font.color.rgb = MUTED
    sub(doc, "A geografia inteira: só estas teclas existem em River Flows in You.")

    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    # ---------- os 4 rituais
    doc.add_heading("🛏️ Os 4 rituais do rio — a mão esquerda inteira em 4 formas", level=1)
    sub(doc, "cada célula vira um ritual: BLOCO → QUEBRADO → LOOP (Fase 1 do plano) · "
             "dedilhado sugerido — porta, não lei: validar com a mão do aluno e anotar no fim")
    rt = doc.add_table(rows=len(RIO_CELULAS) + 1, cols=4)
    rt.style = "Table Grid"
    for ci, h in enumerate(("Célula", "Notas (registro normal)", "Dedos", "Versão funda (seções B)")):
        run = rt.cell(0, ci).paragraphs[0].add_run(h)
        run.bold = True; run.font.size = Pt(9); run.font.color.rgb = GOLD
    for ri, (r, info) in enumerate(RIO_CELULAS.items(), 1):
        vals = (r, " · ".join(info["notas"]), " · ".join(info["dedos"]),
                " · ".join(RIO_GRAVES[r]) if r in RIO_GRAVES else "—")
        for ci, val in enumerate(vals):
            run = rt.cell(ri, ci).paragraphs[0].add_run(val)
            run.font.size = Pt(10)
            run.bold = ci == 0
            if ci == 0:
                run.font.color.rgb = GOLD
            elif ci in (1, 3):
                run.font.name = "Consolas"
    for w, ci in ((0.9, 0), (2.4, 1), (1.1, 2), (2.5, 3)):
        for ri in range(len(RIO_CELULAS) + 1):
            rt.cell(ri, ci).width = Inches(w)
    box(doc, "O loop é a peça",
        "F#m → D → A → E, vi–IV–I–V de Lá maior, do primeiro ao último compasso. Quem "
        "automatiza o loop na mão esquerda já sabe acompanhar a peça inteira — é por isso que "
        "a Fase 1 do plano é só isso.", fill=GOLDB, icon="🌊", tcol=GOLD)

    # ---------- formas de mão da direita
    doc.add_heading("🖐️ As formas de mão da direita", level=1)
    sub(doc, "a peça inteira é feita com estas 4 formas (dedilhado a validar)")
    ft = doc.add_table(rows=len(PADROES_MD) + 1, cols=4)
    ft.style = "Table Grid"
    for ci, h in enumerate(("Forma", "Notas", "Dedos", "O que é")):
        run = ft.cell(0, ci).paragraphs[0].add_run(h)
        run.bold = True; run.font.size = Pt(9); run.font.color.rgb = GOLD
    for ri, (n, ns, dd, obs) in enumerate(PADROES_MD, 1):
        for ci, val in enumerate((n, ns, dd, obs)):
            run = ft.cell(ri, ci).paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            run.bold = ci == 0
            if ci == 0:
                run.font.color.rgb = TEAL
            elif ci == 1:
                run.font.name = "Consolas"
    for w, ci in ((1.15, 0), (2.1, 1), (1.35, 2), (2.3, 3)):
        for ri in range(len(PADROES_MD) + 1):
            ft.cell(ri, ci).width = Inches(w)

    # ---------- a peça compasso a compasso
    doc.add_heading("🌊 A peça, compasso a compasso", level=1)
    sub(doc, "MD = mão direita (azul) · ME = mão esquerda (dourado) · ~ = nota de graça · "
             "VERMELHO = o teto E6 · o rótulo dourado é o rio daquele compasso")
    for nom, i0, i1, tit, desc in SECOES:
        add_secao(doc, nom, i0, i1, tit, desc)
        if nom == 'B':
            box(doc, "m21 — a respiração",
                "A mão direita para no meio do compasso de propósito: meio compasso de G#3 e "
                "silêncio. Não encha: o silêncio é escrito.", icon="🫁")
        if nom == 'B\u2032':
            box(doc, "m38 — a segunda respiração",
                "O mesmo silêncio, uma última vez, antes do convite final.", icon="🫁")
        if nom == 'CODA':
            box(doc, "O final suspenso (m47)",
                "B4 na direita sobre E4 na esquerda, semibreve inteira: o acorde de dominante "
                "que NÃO resolve. Deixe soar. A peça flutua.", icon="🌙")

    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    # ---------- MIDIs
    doc.add_heading("🎧 Os MIDIs de estudo — o professor que toca junto", level=1)
    sub(doc, "gerados pela fábrica na fonte única · as notas de graça viram “crush” de estudo")
    mt = doc.add_table(rows=10, cols=3)
    mt.style = "Table Grid"
    for ci, h in enumerate(("Arquivo", "O que tem", "Pra usar assim")):
        run = mt.cell(0, ci).paragraphs[0].add_run(h)
        run.bold = True; run.font.size = Pt(9); run.font.color.rgb = GOLD
    usos = {
        'esquerda': ("só o rio (mão esquerda)", "tocar a MELODIA por cima, de ouvido"),
        'direita': ("só a melodia (mão direita)", "tocar o RIO por baixo, em bloco ou quebrado"),
        'juntas': ("a peça inteira", "ouvir apontando no mapa · tocar junto · gravar-se por cima"),
    }
    ri = 1
    for mao in ('esquerda', 'direita', 'juntas'):
        for bpm in (60, 70, 78):
            vals = (f"estudo-{mao}-{bpm}.mid",
                    usos[mao][0] + (f" · ♩={bpm}" ),
                    usos[mao][1] if bpm == 60 else
                    ("o degrau do meio" if bpm == 70 else "o andamento da gravação"))
            for ci, val in enumerate(vals):
                run = mt.cell(ri, ci).paragraphs[0].add_run(val)
                run.font.size = Pt(9.5)
                run.bold = ci == 0
                if ci == 0:
                    run.font.name = "Consolas"
                    run.font.color.rgb = TEAL
            ri += 1
    for w, ci in ((2.2, 0), (2.4, 1), (2.3, 2)):
        for rj in range(10):
            mt.cell(rj, ci).width = Inches(w)

    # ---------- o plano
    doc.add_heading("🗺️ O plano — 5 fases, ~9 semanas, 30 min/dia", level=1)
    sub(doc, "rotina diária: 10′ rio (esquerda) · 15′ frase do dia · 5′ revisão do que já estava bom")
    for cod, nom, quando, check in FASES:
        doc.add_heading(f"{cod} · {nom} ({quando})", level=2)
        bullets(doc, check, 9.5)

    # ---------- costuras
    doc.add_heading("🧵 As costuras — onde a peça vira", level=1)
    numlist(doc, [(a + ": ", b) for a, b in COSTURAS])

    # ---------- anotações
    doc.add_heading("✍️ Anotações", level=1)
    sub(doc, "o dedilhado do início é SUGESTÃO — quando a mão do aluno decidir o que é "
             "confortável, é AQUI que o dedilhado validado mora (e depois entra na fábrica)")
    for rot in ("Dedilhado validado — RIO F#m:", "Dedilhado validado — RIO D:",
                "Dedilhado validado — RIO A:", "Dedilhado validado — RIO E:",
                "Dedilhado validado — tema (MD):", "Dedilhado validado — onda e cascata (MD):",
                "Andamentos alcançados:", "Outras descobertas:"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(rot)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(10)
        r2 = p2.add_run("＿" * 52)
        r2.font.color.rgb = MUTED
        r2.font.size = Pt(10)

    # ---------- rodapé
    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(16)
    fr = foot.add_run(
        "Fonte única: river-flow-in-you.mscz (partitura MuseScore 4.40 subida pelo usuário), "
        "parseada por ferramentas/mscx_parse.py — as 854 notas (556 MD + 298 ME) vêm direto da "
        "partitura, nenhuma digitada à mão. Camadas pedagógicas (seções, células, dedilhados, "
        "fases): sugestões da casa, a validar na prática.  ·  Referências: Musicnotes MN0073222 "
        "(original publicado em Lá maior — mesmo tom desta partitura; metrônomo oficial q=65, a "
        "marca ♩=78 é do arquivo). 19 notas de graça e 1 gap de 1/16 (m19) tratados na fábrica.  "
        "·  Regenera: python3 05-river-flows-in-you/trabalho/gera_river_docx.py — os dados moram "
        "em gera_river.py, que lê o .mscz na hora.")
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED

    doc.core_properties.title = "River Flows in You — atril pedagógico (Lá maior)"
    doc.core_properties.author = "Fábrica gera_river.py + gera_river_docx.py"
    doc.save(SAIDAS / "atril-river-flows-in-you.docx")
    return doc


# ── saídas: a pasta ../saidas desta música (a fábrica roda de qualquer lugar) ──
from pathlib import Path as _P
SAIDAS = _P(__file__).resolve().parent.parent / "saidas"
SAIDAS.mkdir(exist_ok=True)

if __name__ == "__main__":
    doc = build()
    d2 = Document(SAIDAS / "atril-river-flows-in-you.docx")
    n_h1 = sum(1 for p in d2.paragraphs if p.style.name == "Heading 1")
    n_h2 = sum(1 for p in d2.paragraphs if p.style.name == "Heading 2")
    n_tbl = len(d2.tables)
    n_rows = sum(len(t.rows) for t in d2.tables)
    print("═" * 74)
    print("DOCX: atril-river-flows-in-you.docx")
    print(f"  H1: {n_h1} · H2: {n_h2} · tabelas: {n_tbl} · linhas de tabela: {n_rows}")
    # 6 seções (2 linhas/compasso) + rituais + formas MD + MIDIs
    esperado = 2 * 47 + len(RIO_CELULAS) + 1 + len(PADROES_MD) + 1 + 9 + 1
    ok = (n_tbl == 9 and n_rows == esperado and n_h1 == 10 and n_h2 == 11
          and len(d2.inline_shapes) == 1)
    print(f"  [{'✓' if ok else '✗'}] estrutura esperada (9 tabelas, {esperado} linhas, "
          f"teclado-guia embutido)")
    raise SystemExit(0 if ok else 1)
