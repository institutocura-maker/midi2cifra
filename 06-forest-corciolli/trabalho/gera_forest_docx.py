#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gera_forest_docx.py — O BARALHO DO ANDARILHO (PoC 6, Forest do Corciolli), versão papel.

Importa gera_forest.py (que re-parseia o .mscz na hora): FONTE ÚNICA de verdade —
nenhuma nota é retranscrita aqui. O que este arquivo acrescenta é a roupa de papel:
cartas A4 pra recortar e jogar o Modo Andarilho na mesa.

BARALHO:
  · MORADAS   (douradas) — onde se mora; cada qual com teclado, ânimo, gramática e cifra
  · TRAVESSIA (×2)       — a carta do caminho (a escala de Si maior)
  · CLAREIRAS (×3)       — a carta do silêncio (1 tempo, 2 tempos, a Grande)
  · GESTOS    (×28)      — as 28 famílias do mapa HTML, com ordem mais comum e durações
  · CARTAS BRANCAS       — as casas vazias da última página: invente a sua morada

Saída: ../saidas/baralho-do-andarilho.docx (+ teclado-forest.png de capa;
os mini-teclados de cada carta são intermediários em trabalho/cartas-png/)
Roda:  python3 06-forest-corciolli/trabalho/gera_forest_docx.py
Requer: pip install python-docx pillow
"""
import os
from collections import Counter

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import gera_forest as gf
from gera_forest import (EPISODIOS, FAM_ORD, GRAMATICA, NOMES_MORADA, MOR_PULSOS,
                         PULSOS, NP, SEG_PULSO, nome, SAIDAS)

# ───────────────────────────────── cores (papel: fundo branco, imprime barato)
INK   = RGBColor(0x21, 0x33, 0x28)   # texto principal (verde-mata escuro)
MUTED = RGBColor(0x5F, 0x74, 0x66)   # rótulos pequenos
GOLD  = RGBColor(0x8A, 0x6D, 0x2F)   # moradas — o ouro em papel
LEAF  = RGBColor(0x2E, 0x6B, 0x3A)   # gestos — o verde-folha
MIST  = RGBColor(0x3E, 0x5F, 0x75)   # clareiras — a névoa
ROAD  = RGBColor(0x4A, 0x5A, 0x6B)   # travessia — a estrada
DARK  = RGBColor(0x0D, 0x24, 0x16)   # títulos

GOLDF = "FFF6E0"                     # fundo das cartas de morada
LEAFF = "F0F9F0"                     # fundo das cartas de gesto
MISTF = "EFF4F8"                     # fundo das cartas de clareira
ROADF = "F4F4F0"                     # fundo das cartas de travessia
BLANF = "FFFFFF"                     # carta branca

PNG_DIR = gf.AQUI / 'cartas-png'
PNG_DIR.mkdir(exist_ok=True)

ACORDE = {11: 'B', 1: 'C#m', 3: 'D#m', 4: 'E', 6: 'F#', 8: 'G#m', 10: 'A# (passagem)'}
DIAT = {0, 2, 4, 5, 7, 9, 11}

EP_POR_PULSO = {}
for _ep in EPISODIOS:
    for _p in range(_ep['ini'], _ep['fim'] + 1):
        EP_POR_PULSO[_p] = _ep


def pretty_dest(d):
    """TRAV/CLAR/nome-de-nota → texto amigável da gramática."""
    if d == 'TRAV':
        return 'travessia'
    if d == 'CLAR':
        return 'clareira'
    return f"{NOMES_MORADA.get(d, (d, ''))[0]} ({d})"


def fmt_dur(x):
    return '%g' % x


# ───────────────────────────────── helpers docx
def shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def cp(cell, label, value, size=8.5, color=INK, italic=False, space=2):
    """parágrafo de carta: rótulo em cor de rótulo + valor em cor de texto."""
    p = cell.paragraphs[0] if (not cell.paragraphs[0].runs and
                               not cell.paragraphs[0].text and
                               len(cell.paragraphs) == 1) else cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space)
    pf.space_after = Pt(0)
    if label:
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = MUTED
    if value:
        r2 = p.add_run(value)
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
        r2.italic = italic
    return p


def tiny_break(doc):
    """quebra de página que quase não ocupa linha."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run()
    r.font.size = Pt(2)
    r.add_break(WD_BREAK.PAGE)


def h1(doc, text, color=DARK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(17)
    r.font.color.rgb = color
    return p


def body(doc, text, size=10.5, color=INK, bold=False, italic=False,
         align=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


# ───────────────────────────────── mini-teclado (PIL)
def _font(sz, bold=True):
    p = ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold
         else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    return ImageFont.truetype(p, sz) if os.path.exists(p) else ImageFont.load_default()


def draw_kb(path, gold=(), green=(), circle=None, scale=False, full=False):
    """mini teclado p/ carta. gold = moradas; green = notas de gesto (ou a escala
    diatônica pálida, se scale=True); circle = pitch circulado (A Casa); full = teclado
    E2–E6 de capa. Teclas destacadas ganham o nome da nota escrito."""
    if full:
        lo, hi = 40, 88                                   # E2..E6
    elif scale:
        lo, hi = 47, 83                                   # B2..B5 (a estrada)
    else:
        ps = list(gold) + list(green) + ([circle] if circle else [])
        lo, hi = min(ps) - 3, max(ps) + 3
        while sum(1 for p in range(lo, hi + 1) if p % 12 in DIAT) < 14:
            lo -= 1
            hi += 1
    whites = [p for p in range(lo, hi + 1) if p % 12 in DIAT]
    W, H, BW, BH, M, y0 = 36, 116, 22, 72, 8, 8
    IMGW = 2 * M + len(whites) * W
    img = Image.new("RGB", (IMGW, y0 + H + 6), "#ffffff")
    d = ImageDraw.Draw(img)
    fk, fb = _font(13), _font(10)

    def center(cx, cy, txt, f, col):
        bb = d.textbbox((0, 0), txt, font=f)
        d.text((cx - (bb[2] - bb[0]) / 2, cy - (bb[3] - bb[1]) / 2), txt, font=f, fill=col)

    gold, green = set(gold), set(green)

    def cols(p):
        if p in gold:
            return "#F5C64B", "#4a3a10"
        if p in green:
            return "#A8DFA0", "#1e4d28"
        if scale and p % 12 in DIAT:
            return "#DDEEDC", None
        return None, None

    x_pos, x = {}, M
    for p in whites:
        fill, tcol = cols(p)
        d.rectangle([x, y0, x + W - 2, y0 + H - 1],
                    fill=fill or "#F7F3E8", outline="#8d9aa8", width=2)
        if tcol:
            center(x + (W - 2) / 2, y0 + H - 20, nome(p), fk, tcol)
        if circle is not None and p == circle:
            cx = x + (W - 2) / 2
            d.ellipse([cx - 11, y0 + 12, cx + 11, y0 + 34], outline="#B03A2E", width=3)
        x_pos[p] = x
        x += W
    for p in range(lo, hi + 1):
        if p % 12 in DIAT:
            continue
        if p - 1 not in x_pos:          # preta de borda (recorte começa nela)
            continue
        bx = x_pos[p - 1] + W - BW // 2 - 1
        fill, tcol = cols(p)
        d.rectangle([bx, y0, bx + BW - 1, y0 + BH - 1],
                    fill=fill or "#23262E", outline="#14171b", width=2)
        if tcol:
            center(bx + BW / 2, y0 + BH - 12, nome(p), fb, tcol)
        x_pos[p] = bx
    img.save(path)


# ───────────────────────────────── as cartas
# linha de carta = (rótulo, valor, estilo) — estilo: 'k' texto, 'i' itálido-cinza,
# 'm' cinza
def carta_morada(m, v):
    """carta dourada: onde se mora."""
    nm = nome(m)
    nm_poet, animo = NOMES_MORADA.get(nm, (nm, 'morada rara — visita rápida'))
    eps = [e for e in EPISODIOS if e['tipo'] == 'morada' and e['ancora'] == m]
    maior = max(e['pulsos'] for e in eps)

    fams = Counter()
    for e in eps:
        for p in range(e['ini'], e['fim'] + 1):
            evs = PULSOS[p - 1]
            if len(evs) >= 2:
                fams[tuple(sorted(x[2] for x in evs))] += 1
    top_fam = fams.most_common(3)

    def fam_txt(f):
        partes = []
        for q, n in Counter(f).items():
            partes.append(nome(q) + ('×2' if n > 1 else ''))
        return '·'.join(partes)

    gestos_txt = ('  ·  '.join(f"{fam_txt(f)} ×{c}" for f, c in top_fam)
                  if top_fam else 'nenhum — só o pedal, nota sobre nota')

    dest = GRAMATICA.get(nm)
    dest_txt = (', '.join(f"{pretty_dest(dd)} ×{c}" for dd, c in dest.most_common(4))
                if dest else '—')

    notas_verdes = {q for f, _ in top_fam for q in f} - {m}
    png = PNG_DIR / f"kb-morada-{nm.replace('#', 's')}.png"
    draw_kb(png, gold={m}, green=notas_verdes, circle=m if nm == 'B4' else None)

    return {
        'kind': 'morada', 'fill': GOLDF, 'tcol': GOLD, 'icon': '🏠',
        'titulo': nm_poet, 'sub': nm, 'png': png,
        'linhas': [
            ('', f"pedal de {nm} · {v} tempos na peça · {len(eps)} visitas · "
                 f"a mais longa: {maior} tempos (~{maior * SEG_PULSO:.0f} s)", 'k'),
            ('', animo, 'i'),
            ('gestos que passam por aqui:  ', gestos_txt, 'k'),
            ('depois daqui o Corciolli foi:  ', dest_txt, 'k'),
            ('cifra por perto:  ',
             f"{ACORDE.get(m % 12, '—')}  (sugestão de ouvido — a mata não marca "
             f"compasso)", 'k'),
            ('pra ouvir:  ', f"morada-{nm}.mid  (a morada em loop, ~30 s)", 'k'),
        ],
    }


def carta_gesto(i, chave, cnt):
    tot = sum(cnt.values())
    realiz = cnt.most_common(1)[0][0]
    ordem = sorted(realiz, key=lambda r: r[0])
    ordem_txt = ' → '.join(nome(r[2]) for r in ordem)
    durs_txt = ' '.join(fmt_dur(r[1]) for r in ordem)

    onde = Counter()
    for pi in range(NP):
        evs = PULSOS[pi]
        if len(evs) >= 2 and tuple(sorted(x[2] for x in evs)) == chave:
            ep = EP_POR_PULSO[pi + 1]
            rot = ('travessia' if ep['tipo'] == 'travessia' else
                   'clareira' if ep['tipo'] == 'clareira' else
                   f"{NOMES_MORADA.get(nome(ep['ancora']), (nome(ep['ancora']), ''))[0]}")
            onde[rot] += 1
    onde_txt = ', '.join(f"{k} ×{c}" for k, c in onde.most_common(3)) if onde else '—'

    png = PNG_DIR / f"kb-gesto-{i:02d}.png"
    draw_kb(png, green=set(chave))

    return {
        'kind': 'gesto', 'fill': LEAFF, 'tcol': LEAF, 'icon': '🍃',
        'titulo': f"gesto nº {i}", 'sub': ' · '.join(nome(p) for p in chave), 'png': png,
        'linhas': [
            ('', f"aparece ×{tot} na peça (variantes únicas são convite, não erro)",
             'i'),
            ('ordem mais comum:  ',
             f"{ordem_txt}   ·   durações (em tempos): {durs_txt}", 'k'),
            ('onde mora:  ', onde_txt, 'k'),
            ('como usar:  ', "solte por cima da morada em que estiver morando — "
             "se for uma das suas, é caminho comprovado", 'k'),
        ],
    }


def carta_travessia():
    png = PNG_DIR / 'kb-travessia.png'
    draw_kb(png, scale=True)
    return {
        'kind': 'travessia', 'fill': ROADF, 'tcol': ROAD, 'icon': '🥾',
        'titulo': 'Travessia', 'sub': 'a estrada', 'png': png,
        'linhas': [
            ('', "caminhe pela escala de Si maior em notas miúdas, até a próxima "
                 "morada", 'k'),
            ('', "é assim que a peça anda entre os campos de pedal: 42 travessias "
                 "gravadas", 'i'),
            ('', "se bater a dúvida, ande devagar — as travessias da gravação "
                 "raramente têm pressa", 'i'),
        ],
    }


def carta_clareira(titulo, tempos, txt, extra):
    return {
        'kind': 'clareira', 'fill': MISTF, 'tcol': MIST, 'icon': '🌫️',
        'titulo': titulo, 'sub': f'{tempos} tempo{"s" if tempos > 1 else ""} de mata',
        'png': None,
        'linhas': [
            ('', txt, 'k'),
            ('', extra, 'i'),
            ('', "não encha — o silêncio aqui é escrito pelo compositor", 'i'),
        ],
    }


def carta_branca():
    return {
        'kind': 'branca', 'fill': BLANF, 'tcol': MUTED, 'icon': '🃏',
        'titulo': 'Carta branca', 'sub': 'a sua morada', 'png': None,
        'linhas': [
            ('', "invente uma morada sua: uma nota, um pedal, um ânimo", 'k'),
            ('', "o modo andarilho livre aceita tudo — contanto que o pedal soe", 'i'),
        ],
    }


def render_card(cell, c):
    cell_shade(cell, c['fill'])
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{c['icon']} {c['titulo']}")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c['tcol']
    r2 = p.add_run(f"  ·  {c['sub']}")
    r2.bold = True
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = INK
    if c['png'] is not None:
        pi = cell.add_paragraph()
        pi.paragraph_format.space_before = Pt(2)
        pi.paragraph_format.space_after = Pt(1)
        pi.add_run().add_picture(str(c['png']), width=Cm(7.5))
    for label, value, st in c['linhas']:
        if st == 'i':
            cp(cell, label, value, color=MUTED, italic=True)
        elif st == 'm':
            cp(cell, label, value, color=MUTED)
        else:
            cp(cell, label, value, color=INK)


def card_page(doc, cards):
    t = doc.add_table(rows=3, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r in t.rows:
        r.height = Cm(8.3)
        r.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for c in r.cells:
            c.width = Cm(8.7)
    n = len(cards)
    for i in range(6):
        cell = t.cell(i // 2, i % 2)
        render_card(cell, cards[i] if i < n else carta_branca())


# ───────────────────────────────── montagem
def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.1)
    sec.left_margin = sec.right_margin = Cm(1.25)

    # ── capa
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run('🌲 FOREST — O BARALHO DO ANDARILHO')
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = DARK
    body(doc, 'Corciolli · A Tranquilidade dos Sons da Natureza (2001) · transcrição '
         'da casa — a 1ª notação da peça', 10, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    body(doc, 'cartas pra recortar · o piano de Si maior que anda por moradas, gestos '
         'e clareiras', 11, INK, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=8)

    kb_capa = SAIDAS / 'teclado-forest.png'
    moradas_notas = {m for m, _ in MOR_PULSOS.most_common(10)}
    gesto_notas = {q for chave, _ in FAM_ORD[:28] for q in chave} - moradas_notas
    draw_kb(kb_capa, gold=moradas_notas, green=gesto_notas, circle=71, full=True)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(kb_capa), width=Cm(17.5))
    body(doc, 'dourado = as moradas (B4 com círculo: A CASA) · verde = as notas dos '
         'gestos · o resto é estrada', 8.5, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)

    body(doc, 'O que é isto', 13, DARK, bold=True, space_after=2)
    for t in [
        'A peça gravada dura 9:20 e é um piano sozinho desenhando campos de pedal por '
        'uma mata de pererecas. A transcrição da casa revelou que ela é uma máquina de '
        'recombinar: 33% do material se repete, e o pedal de Si volta 23 vezes em '
        'pontos distantes. Este baralho destrava essa máquina na sua mesa — sem '
        'deturpar nada: a rota original continua inteira, aqui do lado, nos MIDIs e '
        'no mapa HTML.',
        'A peça não tem partitura publicada; esta transcrição é a primeira. O tom do '
        'mapa é Si maior porque é o tom do piano da casa (meio tom acima da gravação '
        '— piano afinado toca meio tom abaixo do que está escrito aqui). O pulso é de '
        'cerca de 1 segundo; a barra sentida, 2 tempos.',
    ]:
        body(doc, t, 10, INK, space_after=5)

    body(doc, 'Como imprimir e recortar', 13, DARK, bold=True, space_after=2)
    for t in [
        'A4 retrato, frente só (não precisa de verso). As linhas das grades são guias '
        'de corte: 6 cartas por página, ~8,7 × 8,3 cm cada. Papel 90 g/m² já serve; '
        'colar em cartolina 180 g vira baralho de verdade.',
        'São 3 páginas de referência (esta capa, as regras e a gramática) e depois as '
        'cartas: moradas douradas, 2 travessias, 3 clareiras e 28 gestos — as casas '
        'que sobrarem na última página são cartas brancas.',
    ]:
        body(doc, t, 10, INK, space_after=5)

    tiny_break(doc)

    # ── regras
    h1(doc, '🎲 Como se joga o Modo Andarilho')
    body(doc, 'um baralho ou dois dados, uma lei, e a mata por conta das pererecas',
         10.5, MUTED, italic=True, space_after=8)
    passos = [
        ('Escolha uma MORADA (carta dourada)', ' e deixe o pedal soar — re-atacado a '
         'cada tempo, como na gravação. A Casa (B4) é o começo tradicional.'),
        ('Sobre ela, solte GESTOS (cartas verdes)', ' — os da gramática dessa morada '
         'são os seguros: o Corciolli usou-os ali.'),
        ('Quando quiser, atravesse', ': use a carta de travessia e caminhe pela '
         'escala de Si maior em notas miúdas.'),
        ('Chegue na próxima morada', ' — pela gramática (modo seguro) ou livre (a '
         'única lei é o pedal).'),
        ('Respire nas CLAREIRAS', ': um ou dois tempos de mata sozinha. Não encha — '
         'o silêncio é escrito.'),
        ('Pra fechar', ': volte pra Casa (B4) e deixe soar.'),
    ]
    for i, (a, b) in enumerate(passos, 1):
        pp = doc.add_paragraph()
        pp.paragraph_format.space_after = Pt(4)
        pp.paragraph_format.left_indent = Cm(0.5)
        r = pp.add_run(f'{i}. {a}')
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK
        r2 = pp.add_run(b)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = INK

    body(doc, 'Os três modos de tocar', 13, DARK, bold=True, space_after=2)
    for a, b in [
        ('① A Rota', ' — como gravado (fidelidade): toque junto com '
         'rota-original.mid, ou com play-along-com-a-gravacao.mid, que traz os 48 s '
         'de pererecas da intro pra largar junto com o disco.'),
        ('② A Gramática do Corciolli', ' — recombine andando só onde ele andou: a '
         'tabela da próxima página diz pra onde cada morada pode ir.'),
        ('③ O Andarilho livre', ' — baralho na mesa (ou dois dados: um pra morada, '
         'um pra gesto), qualquer carta, a lei do pedal.'),
    ]:
        pp = doc.add_paragraph()
        pp.paragraph_format.space_after = Pt(4)
        r = pp.add_run(a)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK
        r2 = pp.add_run(b)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = INK

    pq = body(doc, 'Deturpar? Não. A peça é 33% repetição e o pedal de Si volta 23 '
              'vezes — ela NASCEU recombinada. A rota original fica aqui, intacta, '
              'pra sempre; o baralho é um modo a mais, declarado como tal.', 10,
              MUTED, italic=True, space_after=8)
    shade(pq, 'F4F9F4')

    body(doc, 'A lei do pedal', 13, DARK, bold=True, space_after=2)
    body(doc, 'A única regra do modo livre: enquanto você mora numa morada, a nota '
         'dela não some — re-atacada a cada tempo, por baixo de qualquer gesto. '
         'Harmonia, pra esta peça, é isso: uma nota que persiste e o que passa por '
         'cima dela. Se o pedal parar, você trocou de morada (ou caiu numa clareira '
         '— tudo bem, contanto que seja de propósito).', 10, INK)

    tiny_break(doc)

    # ── gramática (referência, não se recorta)
    h1(doc, '🧭 A gramática do Corciolli — pra onde cada morada pode ir')
    body(doc, 'depois de cada morada, para onde ele foi na gravação (o modo seguro do '
         'andarilho: ande só onde ele andou)', 10.5, MUTED, italic=True,
         space_after=8)
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    t.autofit = False
    t.rows[0].cells[0].width = Cm(5.2)
    t.rows[0].cells[1].width = Cm(12.2)
    hdr = t.rows[0].cells
    for j, txt in enumerate(('morada', 'destinos seguintes')):
        cell_shade(hdr[j], 'E8F1E8')
        pp = hdr[j].paragraphs[0]
        r = pp.add_run(txt)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = DARK
    for m, _ in MOR_PULSOS.most_common(10):
        nm = nome(m)
        nm_poet, _anim = NOMES_MORADA.get(nm, (nm, ''))
        dest = GRAMATICA.get(nm)
        row = t.add_row()
        row.cells[0].width = Cm(5.2)
        row.cells[1].width = Cm(12.2)
        pp = row.cells[0].paragraphs[0]
        r = pp.add_run(f"{nm_poet} ({nm})")
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = GOLD
        p2 = row.cells[1].paragraphs[0]
        r2 = p2.add_run(', '.join(f"{pretty_dest(d)} ×{v}"
                                  for d, v in dest.most_common(8)) if dest else '—')
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = INK
    body(doc, '', 4)
    body(doc, 'TRAV = travessia (a escala andando) · CLAR = clareira (o silêncio). '
         'Caminhos prontos: caminho-do-andarilho-1/2/3.mid — três caminhadas '
         'sorteadas que só usam passos desta tabela, pra ouvir como soa antes de '
         'tentar.', 9.5, MUTED, italic=True)

    # ── as cartas
    cartas = []
    sem_nome = []
    for m, v in MOR_PULSOS.most_common():
        nm = nome(m)
        if nm in NOMES_MORADA:
            cartas.append(carta_morada(m, v))
        else:
            sem_nome.append(nm)
    n_moradas = len(cartas)
    cartas.append(carta_travessia())
    cartas.append(carta_travessia())
    cartas.append(carta_clareira('Clareira', 1,
                                 'a mata canta sozinha: um tempo de silêncio.',
                                 'na peça: 11 clareiras de 1 tempo'))
    cartas.append(carta_clareira('Clareira dupla', 2,
                                 'dois tempos de mata — a respiração mais funda.',
                                 'na peça: 1 (pouso depois do primeiro campo)'))
    cartas.append(carta_clareira('A Grande Clareira', 4,
                                 'quatro tempos de mata — a pausa maior da peça.',
                                 'na peça: 1 (antes da última volta pra Casa)'))
    for i, (chave, cnt) in enumerate(FAM_ORD[:28], 1):
        cartas.append(carta_gesto(i, chave, cnt))

    n_cartas = len(cartas)
    for k in range(0, n_cartas, 6):
        tiny_break(doc)
        card_page(doc, cartas[k:k + 6])
    n_brancas = (6 - n_cartas % 6) % 6

    out = SAIDAS / 'baralho-do-andarilho.docx'
    doc.save(out)

    # ── auto-checagem de estrutura
    d2 = Document(str(out))
    n_imgs = len(d2.inline_shapes)
    esperado = 1 + n_moradas + 2 + 28            # capa + moradas + travessias + gestos
    assert n_imgs == esperado, f"imagens: {n_imgs} ≠ {esperado}"
    txt = '\n'.join(p.text for p in d2.paragraphs)
    for chave_txt in ('BARALHO DO ANDARILHO', 'Como se joga', 'A lei do pedal'):
        assert chave_txt in txt, f"falta: {chave_txt}"
    n_paginas_carta = (n_cartas + 5) // 6
    assert len(d2.tables) == 1 + n_paginas_carta, "nº de tabelas inesperado"
    corpo = '\n'.join(c.text for tb in d2.tables for row in tb.rows for c in row.cells)
    for m, _ in MOR_PULSOS.most_common(10):
        assert f"morada-{nome(m)}.mid" in corpo, f"carta sem midi: {nome(m)}"
    assert 'A Casa' in corpo and 'Carta branca' in corpo
    assert 'gramática do Corciolli' in txt or 'gramática do Corciolli' in corpo

    print('══════════════════════════════════════════════════════════════════')
    print('FOREST — O BARALHO DO ANDARILHO (papel)')
    print(f'  {n_cartas} cartas: {n_moradas} moradas · 2 travessias · 3 clareiras · '
          f'28 gestos  (+{n_brancas} brancas na última página)')
    if sem_nome:
        print(f'  (moradas raras sem carta: {", ".join(sem_nome)})')
    print(f'  {len(d2.tables)} tabelas (1 gramática + {n_paginas_carta} páginas de '
          f'cartas 2×3) · {n_imgs} imagens embutidas ✓')
    print(f'  → {out}  ({out.stat().st_size // 1024} KB)')
    print('══════════════════════════════════════════════════════════════════')


if __name__ == '__main__':
    build()
