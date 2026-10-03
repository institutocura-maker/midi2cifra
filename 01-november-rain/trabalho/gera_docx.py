#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera mapa-musical-november-rain.docx a partir dos dados de gera_mapa.py
(fonte única de verdade — nenhuma transcrição manual das notas).

Requer:  pip install python-docx pillow
Roda:    python3 01-november-rain/trabalho/gera_docx.py
Saída:   mapa-musical-november-rain.docx (A4, editável, uma tabela por frase)

Organização do documento (pensada pra facilitar versões compactas):
  • Cada FRASE é uma tabela independente (Acorde / Canto-Conta / Notas / Dedos)
    — pra encurtar, é só apagar tabelas ou linhas.
  • Cada seção é um "Heading 1" do Word — dá pra navegar/reorganizar no painel
    de navegação e gerar sumário automático.
  • Notas em vermelho = pico da frase (mesma convenção do HTML).
"""
import gera_mapa as gm
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- cores
INK   = RGBColor(0x22, 0x30, 0x3E)
DARK  = RGBColor(0x1D, 0x2F, 0x4B)
BLUE  = RGBColor(0x1D, 0x5F, 0xA8)
PEAK  = RGBColor(0xC0, 0x39, 0x2B)
MUTED = RGBColor(0x68, 0x78, 0x8A)
CHIP  = RGBColor(0x2C, 0x3E, 0x50)
NOTE  = RGBColor(0x4A, 0x5A, 0x6B)

# ---------------------------------------------------------------- fills (texto puro)
FILLS_TXT = {
    "fill":   "🎹 Piano responde (enquanto a nota presa ainda soa): F#4 · D#4 · E4 · D#4 · C#4 · B3 — "
              "notas compridas, descendo",
    "fill2":  "🎹 Piano responde com uma graça a mais: F#4 · G#4 · F#4 · E4 · D#4 · C#4 · B3",
    "nofill": "🚫 Sem resposta aqui: o \"rain\" segura em cima do vamp E↔F# e a banda cai direto no "
              "refrão — a entrada mais satisfatória da música",
    "bridge": "🌉 Depois do refrão: o F# segura e entra a ponte (\"I know it's hard to keep an open "
              "heart…\") e o 1º solo do Slash",
    "solo3":  "🎸 Sem resposta aqui: o \"rain\" segura e o Slash assume com o 3º solo — e a banda "
              "descola de Si maior para Ré maior…",
    "outro":  "🎤 O motor não para: o Axl já está cantando \"Don't ya think…\" — a seção Final entra "
              "aqui, por cima do solo que ainda roda",
    "fade":   "🔁 " + gm.LOOP_NOTE + ". Cada frase cai em cima do seu compasso — é só deixar rolar "
              "até a música se despedir.",
}

# ---------------------------------------------------------------- seções (título, subtítulo, dados, fills, dicas)
SECTIONS = [
    dict(title="🌧️ Intro A — o tema do piano (compassos 1–15)",
         sub="O começo mais famoso do rock ao piano. Sem letra: a linha do meio conta os TEMPOS "
             "(1 e 2 e 3 e 4 e — conte \"um-e-dois-e…\"), e notas com + tocam juntas "
             "(B3+E4+G#4 = acorde de uma vez). São as notas do disco, sem transposição. A receita: "
             "o baixo desce a escadinha E → D# → C# → B enquanto a direita corre a escala de Si maior.",
         ph=gm.INTRO, fills=gm.INTRO_FILLS,
         tips=[("🎹 A intro é o verso \"de lado\"",
                "Os acordes do tema: E → B/D# → C#m → B — os mesmos acordes do verso (E, C#m, B), "
                "só que com o baixo andando degrau a degrau pela escala (E→D#→C#→B). E as corridas? "
                "A MESMA escala de Si maior dos fills do piano nos versos — lá ela desce "
                "(F#4 D#4 E4 D#4 C#4 B3), aqui ela sobe (E4 F#4 G#4 A#4 B4).")]),
    dict(title="🌩️ Intro B — a entrada da banda (compassos 16–24)",
         sub="Depois do acorde grande, o piano vira base: a direita trava uma moldura e o baixo anda "
             "(E7M → C#m7 → B → B, dois ciclos). Por cima, os sopros fazem a fanfarra que anuncia o verso.",
         ph=gm.GROOVE, fills=gm.GROOVE_FILLS,
         tips=[("✈️ Como estudar a intro",
                "1) A abertura sozinha, até virar músculo: E2 na esquerda, conte 1… 2…, graça D#5 e "
                "caí no E5, segurando. 2) A escadinha do baixo: E → D# → C# → B, um compasso cada, "
                "com a direita parada num acorde simples. 3) As corridas são UMA escala só "
                "(Si maior: E4 F#4 G#4 A#4 B4 — dedos 1-2-3-4-5), devagar, sem pulo. 4) O groove é o "
                "presente: direita parada na moldura, esquerda andando E→C#→B. 5) Monte a sequência "
                "inteira num andamento manso (~60–70 bpm) antes de correr pro tempo do disco.")]),
    dict(title="🌧️ Parte 1 — Verso 1",
         sub="Melodia da voz (mão direita), uma oitava abaixo do Axl. A cifra não mostra ritmo: "
             "você já conhece a música — cante a frase e encaixe as notas.",
         ph=gm.P1, fills=gm.P1_FILLS,
         tips=[("🔑 O \"segredo\" da Parte 1 — a cadência assinatura",
                "Quase toda frase termina no mesmo desenho: C#4 D#4 B3→. E as frases 4, 6 e 8 são "
                "irmãs quase idênticas: C#4 C#4 C#4 E4 D#4 C#4 D#4 B3→.")]),
    dict(title="🎤 Parte 2 — Verso 2",
         sub="Mesma toada da Parte 1, mas o Axl solta mais a voz: tem dois trechos de vocalise "
             "(\"lovers always… go\" e \"today / walking away\"). No mapa está a essência; o enfeite "
             "vem depois.",
         ph=gm.P2, fills=gm.P2_FILLS, tips=[]),
    dict(title="🙌 Parte 2 — Refrão",
         sub="Um molde novo, que se repete 4 vezes: D#4 C#4 C#4 C#4 B3 B3 C#4 + a palavra final. "
             "Acordes do disco: D#m → E → B/D# (linhas 1 e 3) e D#m → E → C#m (linha 2); a cifra "
             "comum simplifica pra E → F# por linha.",
         ph=gm.CH, fills=gm.CH_FILLS,
         tips=[("🧲 Por que o refrão \"flutua\" (e nunca pousa)",
                "Os acordes ficam rodando em D#m → E → B/D# → C#m → F# — e nunca chegam no B "
                "(a tônica, o \"casa\" da música). Tudo fica suspenso, como uma pergunta no ar. "
                "No piano: comece com a cifra simples (E → F#) e, quando estiver firme, troque "
                "pelos acordes reais.")]),
    dict(title="🌉 Ponte — \"I know it's hard to keep an open heart…\"",
         sub="Transparência: o nosso MIDI não tem a voz nesta seção (a pista vocal fica muda) e não "
             "há transcrição confiável da linha cantada — o mapa traz a cama de acordes (conferida "
             "em 4 fontes) e a letra. O canto do Axl é baixinho, quase falado — cante por cima. "
             "É a mesma viagem harmônica do refrão, com a última linha terminando em F# seco.",
         ph=gm.PONTE, fills=gm.PONTE_FILLS,
         tips=[("🧩 A ordem da seção (pra não se perder)",
                "Refrão 1 → ponte (4 linhas, ~2 compassos cada) → solo 1 (16 compassos) → refrão 2 "
                "(\"Sometimes I need some time…\"). No nosso MIDI o transcritor pulou a cama da ponte "
                "e deixou o solo cobrir os 16 compassos — no disco são ~8 compassos de ponte cantada "
                "e depois o solo. E os acordes do solo são os do verso de novo (E, C#m, B).")]),
    dict(title="🎸 Solo 1 — a resposta do Slash (compassos 77–92)",
         sub="Nota a nota do nosso MIDI, no registro real da guitarra — a essência do solo "
             "(repetições rápidas viram notas presas; os enfeites ficam por sua conta). Cama: "
             "E → C#m → B, com o baixo passeando (B/D#, B/F#). Se quiser, toque tudo uma oitava "
             "abaixo: o desenho é o que importa.",
         ph=gm.SOLO1, fills=gm.SOLO1_FILLS,
         tips=[("🎸 A arquitetura do solo (o mapa do arco)",
                "Conversa (o vai-e-vem G#↔F#, c. 77–78) → repetição um degrau acima (81–82) → "
                "o grito (D#5, c. 83) → o canto (E5, c. 85) → O TETO (B5!, c. 87) → a grande escada "
                "descendo (c. 88) → as terças em queda (c. 90) → a aterrissagem (de volta ao D#4 "
                "onde começou, c. 91) → a entrega pro refrão 2 (c. 92)."),
               ("🔁 Bônus: o 2º refrão é uma variação",
                "Na segunda volta Axl canta \"Sometimes I need some time on my own\" — mesmo molde, "
                "abertura diferente: Some-(B3) times(C#4) I(B3) need(E4) some(D#4) time(B3→) "
                "on(C#4) my(D#4) own(C#4→).")]),
    dict(title="🎸 Solo 2 — o chorão (compassos 100–109)",
         sub="Logo depois do refrão 2 (\"Sometimes I need some time…\"), que termina segurando o F# "
             "(a pergunta no ar) — e o Slash responde com uma escada de lançamento inteira. Cama: "
             "F# → E → C#m → B → B/D# (duas voltas). O momento famoso: um B5 segurado por mais de "
             "4 tempos — na guitarra é o bend chorado; no piano, pedale e deixe cantar.",
         ph=gm.SOLO2, fills=gm.SOLO2_FILLS,
         tips=[("🎸 Os três irmãos — Solo 1, Solo 2 e Solo 3",
                "O Solo 1 conversa: vai-e-vem, repetições, um grito e volta pra casa. O Solo 2 chora: "
                "sobe uma escada inteira, salta dois andares e segura a nota mais longa da música. "
                "O Solo 3 viaja: em vez de uma frase, é a música inteira — cama, silêncio, um mundo "
                "novo e o teto absoluto (C#6!). Os três têm a escala de Si maior nas veias, e cada "
                "um termina na porta da próxima seção. Aprende um, metade do outro já está na mão.")]),
    dict(title="🌅 Parte 3 — Verso 3",
         sub="O verso da despedida — e um presente pra quem estudou a Parte 1: quase tudo aqui é "
             "reciclagem das melodias que você já sabe.",
         ph=gm.P3, fills=gm.P3_FILLS,
         tips=[("♻️ O Verso 3 é um \"greatest hits\" da Parte 1",
                "\"'Cause nothin' lasts forever\" e \"Even cold November rain\" são literalmente as "
                "frases 5 e 8 da Parte 1 (nota por nota). \"We still can find a way\" é irmã das "
                "frases 4/6/8, e a abertura \"And when your fears subside\" recicla o pulo do "
                "\"We've been through this\".")]),
    dict(title="🎸 Solo 3 — o gigante (compassos 122–151)",
         sub="O solo mais longo da música — 30 compassos — e o único que é uma VIAGEM: a cama "
             "clássica de Si maior, um mergulho com compasso de silêncio TOTAL, o groove G | A que "
             "muda o mundo (e a velocidade: 80 → 105 bpm!) e o clímax com a nota mais aguda de toda "
             "a música: C#6. Lead no registro real; linha grave = cama (baixo). Se doer a mão, uma "
             "oitava abaixo — o desenho é o que importa.",
         ph=gm.SOLO3, fills=gm.SOLO3_FILLS,
         tips=[("🎸 A arquitetura do gigante — 5 fases e uma porta",
                "1) A cama (c. 124–131): B | F#/A# | G#m | F# | E | F# ×2 — a escadinha "
                "B→A#→G#→F#→E descendo nos graves, um acorde por degrau. 2) O mergulho (c. 132–133): "
                "o último eco de Si… e um compasso de silêncio TOTAL — não toca nada, nem bateria. "
                "3) O novo mundo (c. 134–141): o piano assume o groove G | A | B — o mesmo do final! "
                "O G natural é a senha do mundo novo, e a música acelera de 80 pra 105 bpm "
                "exatamente aqui. 4) O clímax (c. 142–149): sobre a harmonia do final comprimida "
                "(Bm | F# | E–D | G–A), a lead grita B5 → A#5 → C#6!, flutua F#–G–F#–E e martela o "
                "uníssono B+D — e depois repete TUDO, nota por nota. 5) O motor (c. 150–151): o trem "
                "F#–G–F#–E / C#–D–C#–B que não para mais — enquanto o Axl já canta "
                "\"Don't ya think…\" por cima. Os outros solos voltam pra casa; este abre uma porta "
                "e se dissolve no final.")]),
    dict(title="🌟 Parte 3 — Final em Ré maior (\"everybody needs somebody\", compassos 150–169)",
         sub="Aqui a música sobe de Si maior para Ré maior (uma 3ª menor acima) — por isso o "
             "encerramento soa \"iluminado\". A voz já canta na região do dó central (D4–E4): estas "
             "notas são as da gravação, sem transposição nenhuma.",
         ph=gm.OUT, fills=gm.OUT_FILLS,
         tips=[("🌟 A modulação do final — o \"sol depois da chuva\"",
                "A música inteira mora em Si maior… e o final sobe pra Ré maior — mas não do nada: "
                "o Solo 3 já preparou o terreno com o groove G | A e o G natural (desde o c. 134!). "
                "O groove é um loop de 5 compassos: Bm | Bm | Bm–F# | E–D | G–A — o E maior no meio "
                "é o tempero, e o G–A no fim é a mola que joga tudo de volta pro Bm. Na esquerda: "
                "um baixo por acorde (B2 · F#2 · E2 · D2 · G2 · A2), oitava 2, bem fraquinho."),
               ("✈️ Como estudar a Parte 3",
                "1) Comece pelo que é de graça: \"'Cause nothin' lasts forever / Even cold November "
                "rain\" você JÁ sabe (frases 5 e 8 da Parte 1, nota por nota). 2) As duas frases "
                "novas: \"And when your fears subside\" (o pulo B3→G#4) e o par \"I know that you "
                "can love me / When there's no one left to blame\" (o mergulho no G#3 e a subida no "
                "\"no\"). 3) O final é um loop: aprenda as 5 frases, ligue cada uma ao seu acorde do "
                "loop (Bm · Bm · Bm–F# · E–D · G–A) e repita 4× — fica bonito subir uma oitava na "
                "última repetição. 4) Grave-se tocando junto com o trecho final do disco.")]),
]

LH_ROWS = [
    ("Acorde", "Simplezinho (1ª etapa)", "Com corpo (2ª etapa)", "Onde aparece"),
    ("E7M", "E2 + E3", "E2 + G#3+Si3+Ré#4", "versos (o Ré# é a 7ª maior — o som \"docinho\")"),
    ("C#m7", "C#2 + C#3", "C#2 + Sol#2+C#3+Mi3", "versos"),
    ("F# / F#7", "F#2 + F#3", "F#2 + Lá#3+Dó#4", "\"for-E-ver\", \"If we could take the time\", \"We still can find a way\""),
    ("B", "B2 (ou B1+B2)", "B2 + Fá#3+Si3", "fim de frase (cadência) e intro (o acorde grande do c. 15)"),
    ("B/D#", "D#2", "D#2 + Fá#3+Si3", "intro (o 2º degrau da escadinha) e refrão"),
    ("D#m", "D#2 + D#3", "D#2 + Fá#3+Lá#3", "refrão e ponte"),
    ("Final (Ré maior)", "B2 / F#2 / E2 / D2 / G2 / A2", "um por acorde do loop", "o loop Bm | Bm | Bm–F# | E–D | G–A"),
]

COSTURAS = [
    "1) Refrão 1 → Ponte: o refrão morre no F# suspenso (pergunta no ar) — a ponte responde "
    "baixinho, descendo pro D#m. A queda de volume mais dramática da música.",
    "2) Ponte → Solo 1: a última linha termina em F# SECO — e o Slash decola no vai-e-vem G#↔F#.",
    "3) Refrão 2 → Solo 2: de novo o F# suspenso — e a resposta é a escada C#4→C#5.",
    "4) Solo 2 → Verso 3: o solo aterrissa no C#5 e o \"And when your fears subside\" entra no B3 — "
    "uma nota de distância.",
    "5) Verso 3 → Solo 3: o \"Even cold November rain\" segura no B3→, a guitarra responde gritando "
    "(F#5 → B4)… e depois da cama clássica vem o mergulho: um compasso de silêncio TOTAL — e o "
    "piano assume o groove G | A (a música acelera pra 105 bpm aqui!).",
    "6) Solo 3 → Final: quando o motor F#–G–F#–E começar a rodar, prepare a mão: \"Don't ya "
    "think…\" entra no D4 — a primeira nota da região nova (Ré maior!) — por cima do solo, que "
    "continua rodando até o fade.",
]

ROTEIRO = [
    "Antes de tudo: ouça o disco com o guia aberto, acompanhando as 12 seções — marque onde cada "
    "uma entra (a letra é a deixa).",
    "Ensaie só as costuras: a última frase de cada seção + a primeira da seguinte (a lista das 6 "
    "acima). A mão tropeça aí, não no meio.",
    "Nos solos, toque a cama (ou cante junto — também conta!).",
    "Se perder, NÃO pare: segura o acorde da seção (quase sempre E, C#m ou B) e espera a próxima "
    "entrada — a música sempre volta pra casa.",
    "Duas ou três rodadas assim, depois grave-se e compare com o disco. Aí é só curtir. 🎹",
]

# ---------------------------------------------------------------- teclado (PIL)
def draw_keyboard(path):
    from PIL import Image, ImageDraw, ImageFont
    WHITEK = ["C", "D", "E", "F", "G", "A", "B"]
    BLACKK = {"C": "C#", "D": "D#", "F": "F#", "G": "G#", "A": "A#"}
    mel = {"F#3", "G#3", "A3", "B3", "C#4", "D4", "D#4", "E4", "F#4", "G#4", "A#4", "B4",
           "C#5", "D#5", "E5", "F#5", "G#5", "A#5", "B5", "C#6"}
    lh = {"E2", "C#2", "D#2", "F#2", "B2", "D2", "G2", "A2", "A#2", "G#2"}
    W, H, BW, BH, M = 64, 280, 40, 172, 12
    n_white = 29
    img = Image.new("RGB", (n_white * W + 2 * M, M + H + 64), "#eef2f7")
    d = ImageDraw.Draw(img)

    def font(sz, bold=True):
        for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold
                  else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",):
            import os
            if os.path.exists(p):
                return ImageFont.truetype(p, sz)
        return ImageFont.load_default()

    f_key, f_oct, f_leg = font(15), font(13, bold=False), font(15)

    def center_text(cx, cy, txt, f, col):
        bb = d.textbbox((0, 0), txt, font=f)
        d.text((cx - (bb[2] - bb[0]) / 2, cy - (bb[3] - bb[1]) / 2), txt, font=f, fill=col)

    y0 = M
    x = M
    c4x = None
    for octv in (2, 3, 4, 5):
        for wname in WHITEK:
            label = f"{wname}{octv}"
            if label in mel:
                fill, tcol = "#cfe3f8", "#1d5c9e"
            elif label in lh:
                fill, tcol = "#f9e4c8", "#b06a1a"
            else:
                fill, tcol = "#ffffff", "#5a6b7d"
            d.rectangle([x, y0, x + W - 1, y0 + H - 1], fill=fill, outline="#8d9aa8", width=2)
            if label in mel or label in lh:
                center_text(x + W / 2, y0 + H - 52, label, f_key, tcol)
            if wname == "C":
                center_text(x + W / 2, y0 + H - 24, label, f_oct, "#7d8b99")
            if label == "C4":
                c4x = x + W / 2
                d.ellipse([c4x - 7, y0 + 96, c4x + 7, y0 + 110], fill="#e67e22")
                center_text(c4x, y0 + 84, "dó central", font(12, bold=False), "#e67e22")
            x += W
    # C6 final
    d.rectangle([x, y0, x + W - 1, y0 + H - 1], fill="#ffffff", outline="#8d9aa8", width=2)
    center_text(x + W / 2, y0 + H - 24, "C6", f_oct, "#7d8b99")
    # pretas
    x = M
    for octv in (2, 3, 4, 5):
        for wname in WHITEK:
            if wname in BLACKK:
                bl = f"{BLACKK[wname]}{octv}"
                bx = x + W - BW // 2
                if bl in mel:
                    bfill, btxt = "#2f6fb5", bl
                elif bl in lh:
                    bfill, btxt = "#d08a3e", bl
                else:
                    bfill, btxt = "#2b3038", None
                d.rectangle([bx, y0, bx + BW - 1, y0 + BH - 1], fill=bfill, outline="#14171b", width=2)
                if btxt:
                    center_text(bx + BW / 2, y0 + BH - 18, btxt, font(11), "#ffffff")
            x += W
    bx = x + W - BW // 2
    d.rectangle([bx, y0, bx + BW - 1, y0 + BH - 1], fill="#2f6fb5", outline="#14171b", width=2)
    center_text(bx + BW / 2, y0 + BH - 18, "C#6", font(11), "#ffffff")
    # legenda
    ly = y0 + H + 18
    d.rectangle([M, ly, M + 26, ly + 16], fill="#cfe3f8", outline="#8d9aa8")
    d.text((M + 34, ly), "melodia (mão direita)", font=f_leg, fill="#33465a")
    x2 = M + 34 + d.textlength("melodia (mão direita)", font=f_leg) + 26
    d.rectangle([x2, ly, x2 + 26, ly + 16], fill="#f9e4c8", outline="#8d9aa8")
    d.text((x2 + 34, ly), "baixos (mão esquerda)", font=f_leg, fill="#33465a")
    x3 = x2 + 34 + d.textlength("baixos (mão esquerda)", font=f_leg) + 26
    d.ellipse([x3 + 6, ly + 2, x3 + 20, ly + 16], fill="#e67e22")
    d.text((x3 + 30, ly), "C4 = dó central", font=f_leg, fill="#33465a")
    img.save(path)

# ---------------------------------------------------------------- helpers docx
def shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def tip_box(doc, title, text):
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_before = Pt(8)
    p1.paragraph_format.space_after = Pt(0)
    r = p1.add_run("💡 " + title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x8A, 0x5D, 0x0B)
    shade(p1, "FFF8EC")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)
    r2.font.color.rgb = NOTE
    shade(p2, "FFF8EC")


def add_phrase(doc, p, fallback_label):
    label = p.get("label") or fallback_label
    pre = p.get("pre")
    if pre == "PREVAMP":
        tip_box(doc, "O vamp antes do refrão", gm.VAMP_NOTE)
    elif pre == "LOOP":
        tip_box(doc, "O loop do final", gm.LOOP_NOTE + ". Cada frase cai em cima do seu compasso.")
    elif pre == "MOLDURA":
        tip_box(doc, "A moldura parada — o truque do groove", gm.MOLDURA_NOTE)

    lp = doc.add_paragraph()
    lp.paragraph_format.space_before = Pt(10)
    lp.paragraph_format.space_after = Pt(2)
    lr = lp.add_run("◆ " + label)
    lr.bold = True
    lr.font.size = Pt(9.5)
    lr.font.color.rgb = MUTED

    ncols = len(p["syls"])
    has_fin = bool(p.get("fingers"))
    nrows = 4 if has_fin else 3
    tbl = doc.add_table(rows=nrows, cols=ncols + 1)
    tbl.style = "Table Grid"
    chords_at = {i: c for c, i in p["chords"]}
    chord_row = ["Acorde"] + [chords_at.get(i, "") for i in range(ncols)]
    syl_row = ["Canto/Conta"] + list(p["syls"])
    notes_row = ["Notas (MD)"] + [" › ".join(ns) for ns in p["notes"]]
    rows = [chord_row, syl_row, notes_row]
    if has_fin:
        fins = list(p["fingers"]) + [""] * (ncols - len(p["fingers"]))
        rows.append(["Dedos"] + fins[:ncols])
    peak = p.get("peak")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(val)
            if ci == 0:
                run.font.size = Pt(8)
                run.font.color.rgb = MUTED
                run.bold = True
            elif ri == 0:      # acorde
                run.bold = True
                run.italic = True
                run.font.size = Pt(10)
                run.font.color.rgb = CHIP
            elif ri == 1:      # canto/conta
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = INK
            elif ri == 2:      # notas
                run.font.name = "Consolas"
                run.bold = True
                run.font.size = Pt(9.5)
                if ci - 1 == peak:
                    run.font.color.rgb = PEAK
                else:
                    run.font.color.rgb = BLUE
            else:              # dedos
                run.font.size = Pt(8.5)
                run.font.color.rgb = MUTED
    # larguras
    tbl.autofit = False
    total = 6.7
    first = 1.05
    wcol = (total - first) / ncols
    for ri in range(nrows):
        tbl.cell(ri, 0).width = Inches(first)
        for ci in range(1, ncols + 1):
            tbl.cell(ri, ci).width = Inches(wcol)
    if p.get("note"):
        np_ = doc.add_paragraph()
        np_.paragraph_format.space_before = Pt(2)
        np_.paragraph_format.space_after = Pt(4)
        nr = np_.add_run("💡 " + p["note"])
        nr.italic = True
        nr.font.size = Pt(9)
        nr.font.color.rgb = NOTE


def add_fill(doc, kind):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("↳ " + FILLS_TXT[kind])
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED


# ---------------------------------------------------------------- montagem
def build():
    doc = Document()
    # página A4 + margens
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    st = doc.styles["Normal"]
    st.font.name = "Segoe UI"
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Segoe UI"
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = DARK

    # ---------- cabeçalho
    t = doc.add_heading("🎹 Mapa Musical — November Rain", level=0)
    sub = doc.add_paragraph()
    r = sub.add_run("Guns N' Roses · Use Your Illusion I — O mapa completo: os 3 solos + guia de "
                    "montagem, do primeiro acorde ao fade")
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED
    badges = doc.add_paragraph()
    rb = badges.add_run("Tom: Si maior (final em Ré maior!)   ·   Cifras: E7M · C#m7 · B/D# · B · "
                        "F# · D#m · Bm…   ·   Voz 1 oitava abaixo; intro, solos e final no registro "
                        "real   ·   Andamento: ~80 bpm (acelera pra ~105 bpm no final)")
    rb.font.size = Pt(9)
    rb.font.color.rgb = MUTED

    # ---------- como ler
    doc.add_heading("📖 Como ler o mapa", level=1)
    for b in [
        "Números = oitavas. C4 é o dó central. G#3 fica logo abaixo dele; D#4 logo acima.",
        "→ nota longa: toca e deixa soar. · = silêncio. › liga notas da mesma sílaba (melisma).",
        "Na intro não tem letra: a linha do meio são os tempos (1 e 2 e 3 e 4 e — conte "
        "\"um-e-dois-e…\"), e notas com + tocam juntas: B3+E4+G#4 = acorde de uma vez.",
        "Notas em VERMELHO = pico da frase (onde a voz abre). Números cinza = dedilhado sugerido "
        "(mão direita) — ajuste ao que for confortável.",
        "Notas repetidas são o Axl re-atacando a mesma nota — toque uma vez e deixe cantar.",
        "Cifra de violão × piano: muita cifra mostra F–Dm–C com \"afine meio tom abaixo\". No piano, "
        "toque direto E7M, C#m7, B, F# — é o som real do disco.",
    ]:
        p = doc.add_paragraph(b, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(10)

    # ---------- teclado
    doc.add_heading("🎼 Onde estão as notas (teclado-guia)", level=1)
    try:
        draw_keyboard("/tmp/teclado.png")
        doc.add_picture("/tmp/teclado.png", width=Inches(6.7))
        doc.paragraphs[-1].alignment = 1
    except Exception as e:
        p = doc.add_paragraph(f"(Teclado-guia indisponível nesta geração: {e})")
        p.runs[0].font.color.rgb = MUTED
    kp = doc.add_paragraph()
    kr = kp.add_run("A Parte 1 inteira mora em 8 teclas (F#3 → G#4). Partes 1–3 juntas: 11 teclas — "
                    "a Parte 2 acrescentou só a B4, e o final trouxe a A3 e a D4. A intro estica o "
                    "mapa pra cima (A#4, C#5, D#5, E5, F#5). E os solos do Slash vão além: G#5, A#5, "
                    "B5… e o Solo 3 coroa no C#6 — a nota mais aguda da música inteira, o teto "
                    "absoluto do mapa (no Solo 2, o B5 ainda é a nota mais longa). Repare como o "
                    "mapa \"anda\" pra direita no teclado conforme a música avança. 😉")
    kr.font.size = Pt(9.5)
    kr.font.color.rgb = NOTE

    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    # ---------- seções musicais
    for si, s in enumerate(SECTIONS, 1):
        doc.add_heading(s["title"], level=1)
        ps = doc.add_paragraph()
        pr = ps.add_run(s["sub"])
        pr.italic = True
        pr.font.size = Pt(9.5)
        pr.font.color.rgb = NOTE
        for k, p in enumerate(s["ph"]):
            add_phrase(doc, p, f"Frase {k + 1}")
            if k in s["fills"]:
                add_fill(doc, s["fills"][k])
        for tt, tx in s.get("tips", []):
            tip_box(doc, tt, tx)

    # ---------- mão esquerda
    doc.add_heading("🖐️ Mão esquerda (acompanhamento)", level=1)
    p = doc.add_paragraph("Comece simples — só o baixo, bem fraquinho:")
    p.runs[0].font.size = Pt(10)
    lh = doc.add_table(rows=len(LH_ROWS), cols=4)
    lh.style = "Table Grid"
    widths = [Inches(1.15), Inches(1.7), Inches(1.85), Inches(2.0)]
    for ri, row in enumerate(LH_ROWS):
        for ci, val in enumerate(row):
            cell = lh.cell(ri, ci)
            cell.width = widths[ci]
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(val)
            run.font.size = Pt(8.5 if ri else 9)
            if ri == 0:
                run.bold = True
                run.font.color.rgb = CHIP
            elif ci == 0:
                run.bold = True
    up = doc.add_paragraph()
    ur = up.add_run("💡 Upgrade: no refrão, o disco alterna D#m → E → B/D# (ou C#m) dentro de cada "
                    "linha; no fim do verso 2 (\"So if you want to love me…\") é um vamp E ↔ F# "
                    "compasso a compasso até estourar no refrão. Na intro: o baixo desce a escadinha "
                    "E → D# → C# → B (tema) e depois pulsa E → C# → B (groove, 2×). E no Solo 3, a "
                    "escadinha desce de B: B → A# → G# → F# → E.")
    ur.italic = True
    ur.font.size = Pt(9)
    ur.font.color.rgb = NOTE

    # ---------- guia de montagem
    doc.add_heading("🗺️ Guia de montagem — a música inteira, de ponta a ponta", level=1)
    gp = doc.add_paragraph()
    gr = gp.add_run("169 compassos, ~9 minutos — e você tem TODAS as notas: é só seguir a fila. "
                    "Os números de compasso são os do nosso MIDI. A letra entre aspas é a deixa: "
                    "quando ouvir ela chegando, a seção é aquela.")
    gr.font.size = Pt(9.5)
    gr.font.color.rgb = NOTE
    gt = doc.add_table(rows=len(gm.GUIDE) + 1, cols=6)
    gt.style = "Table Grid"
    gw = [Inches(0.3), Inches(1.5), Inches(1.55), Inches(1.85), Inches(0.9), Inches(0.6)]
    heads = ["#", "Seção", "Como reconhece", "Acordes", "No mapa", "Comp."]
    for ci, h in enumerate(heads):
        cell = gt.cell(0, ci)
        cell.width = gw[ci]
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = CHIP
    for ri, (n, como, ac, mp, comp) in enumerate(gm.GUIDE, 1):
        vals = [str(ri), n, como, ac, mp, comp]
        for ci, val in enumerate(vals):
            cell = gt.cell(ri, ci)
            cell.width = gw[ci]
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(8.5)
            if ci == 1:
                run.bold = True
    gn = doc.add_paragraph()
    gnr = gn.add_run(gm.GUIDE_NOTE)
    gnr.italic = True
    gnr.font.size = Pt(8.5)
    gnr.font.color.rgb = MUTED

    tip_box(doc, "🎸 E os solos? (como tocar a música inteira sem tocar os solos)",
            "Nos solos 1 e 2 a cama é a família do verso (E, C#m, B, F#) — toque como se fosse um "
            "verso sem letra: baixo nos tempos, moldura na direita, e deixe o disco brilhar por "
            "cima. No solo 3, a cama é a escadinha nos graves (B → A# → G# → F# → E) e depois o "
            "groove G | A — o mesmo do final. E os três solos estão transcritos no mapa, se um dia "
            "quiser brilhar você mesmo. Lembre: no disco o piano NUNCA para — tocar a cama já é "
            "tocar a música.")

    tip_box(doc, "🧵 As 6 costuras (é onde todo mundo se perde)", " ".join(COSTURAS))

    p = doc.add_paragraph()
    pr = p.add_run("✈️ Roteiro pra primeira rodada inteira")
    pr.bold = True
    pr.font.size = Pt(11)
    for i, item in enumerate(ROTEIRO, 1):
        ip = doc.add_paragraph()
        ip.paragraph_format.space_after = Pt(2)
        ir = ip.add_run(f"{i}. {item}")
        ir.font.size = Pt(10)

    tip_box(doc, "🏆 Status do projeto",
            "Toda a estrutura da música está mapeada: intro, versos 1–3, refrões (com a variação do "
            "2º), ponte (cama + letra), o final em Ré maior — e, de bônus, os TRÊS solos nota a "
            "nota, do vai-e-vem do Solo 1 ao C#6 do gigante. Sobrou pra um dia, se a curiosidade "
            "apertar: só a linha cantada da ponte (que não existe no nosso MIDI).")

    ft = doc.add_paragraph()
    fr = ft.add_run("Mapa conferido nota a nota com a gravação original (MIDI da performance). "
                    "Versos/refrão na região F#3–B4 (uma oitava abaixo da voz); intro, solos e "
                    "final no registro real — a intro abre no E5, os solos 1 e 2 coroam no B5 e o "
                    "solo 3 chega ao C#6; o final mora no D4–E4, já em Ré maior. Versão DOCX gerada "
                    "por gera_docx.py a partir de gera_mapa.py.")
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED

    doc.save(SAIDAS / "mapa-musical-november-rain.docx")
    return doc


# ── saídas: a pasta ../saidas desta música (a fábrica roda de qualquer lugar) ──
from pathlib import Path as _P
SAIDAS = _P(__file__).resolve().parent.parent / "saidas"
SAIDAS.mkdir(exist_ok=True)

if __name__ == "__main__":
    d = build()
    n_tbl = len(d.tables)
    n_par = len(d.paragraphs)
    print(f"[OK] mapa-musical-november-rain.docx gerado — {n_tbl} tabelas, {n_par} parágrafos")
